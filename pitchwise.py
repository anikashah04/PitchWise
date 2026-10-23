import os
import re
import asyncio
import threading
from pathlib import Path
from datetime import datetime

from crewai import Agent, Task, Crew, Process, LLM
from docx import Document
from docx.shared import Inches as DocxInches, Pt as DocxPt, RGBColor as DocxRGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor


DEFAULT_MODEL = "ollama_chat/gemma4:31b"
OUTPUT_DIR = Path(os.environ.get("PITCHWISE_OUTPUT_DIR", ".")).resolve()

# Rules appended to every task prompt (keeps agents consistent and concise)
COMMON_RULES = """
General rules:
- Plain-text math only: write "x" and "approx.", never LaTeX or $...$ symbols. Use Rs, lakh and crore as in the pitch (1 crore = 100 lakh).
- Do not attribute a strategy or fact to the pitch (for example "land and expand", "freemium", a sales pipeline) unless the pitch states it.
- Any competitor, tool or market fact that comes from your general knowledge must be marked "(unverified)".
- Be concise and do not repeat a point that already appears in another section.
"""

# Presentation palette
NAVY = RGBColor(22, 37, 66)
BLUE = RGBColor(48, 95, 170)
TEAL = RGBColor(25, 139, 145)
INK = RGBColor(45, 55, 72)
MUTED = RGBColor(103, 114, 130)
PALE = RGBColor(242, 246, 251)
WHITE = RGBColor(255, 255, 255)
GOLD = RGBColor(218, 164, 65)


def get_llm(api_key=None, model=None):
    """Create the CrewAI LLM client. Ollama Cloud by default; any other
    provider works by setting PITCHWISE_MODEL, e.g. gemini/..., groq/..., openai/...
    (their keys are read from GEMINI_API_KEY, GROQ_API_KEY, OPENAI_API_KEY)."""
    model = model or os.environ.get("PITCHWISE_MODEL", DEFAULT_MODEL)

    if model.startswith("ollama"):
        api_key = api_key or os.environ.get("OLLAMA_API_KEY")
        if not api_key:
            raise ValueError(
                "OLLAMA_API_KEY is not set. Add it to Colab Secrets and load it "
                "into os.environ before calling run_pitchwise()."
            )
        return LLM(
            model=model,
            api_key=api_key,
            base_url="https://ollama.com",
            temperature=0.25,
            max_tokens=3500,
        )

    kwargs = {"model": model, "temperature": 0.25, "max_tokens": 3500}
    if api_key:
        kwargs["api_key"] = api_key
    return LLM(**kwargs)


def _safe_filename(name):
    name = re.sub(r"[^A-Za-z0-9._-]+", "_", str(name or "Startup"))
    return name.strip("._") or "Startup"


def _clean_latex(text):
    """Replace stray LaTeX math that models sometimes emit with plain text."""
    text = str(text or "")
    replacements = {
        r"$\times$": "x", r"\times": "x",
        r"$\approx$": "approx.", r"\approx": "approx.",
        r"$\geq$": ">=", r"\geq": ">=",
        r"$\leq$": "<=", r"\leq": "<=",
        r"$\rightarrow$": "->", r"\rightarrow": "->",
    }
    for src, dst in replacements.items():
        text = text.replace(src, dst)
    return text


def _strip_markdown(text):
    text = re.sub(r"```.*?```", "", str(text or ""), flags=re.S)
    text = re.sub(r"^\s{0,3}#{1,6}\s*", "", text, flags=re.M)
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"(?<!\*)\*(?!\s)(.*?)(?<!\s)\*(?!\*)", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text.strip()


def _extract_startup_name(pitch):
    match = re.search(r"(?im)^\s*Startup\s*:\s*(.+?)\s*$", pitch or "")
    if match:
        return match.group(1).strip()
    match = re.search(r"(?im)^\s*(?:Company|Business|Venture)\s*:\s*(.+?)\s*$", pitch or "")
    return match.group(1).strip() if match else "Startup Assessment"


def _extract_verdict(report_text):
    """Find the verdict from a 'Verdict: X' line, not from random uses of the word 'go'."""
    match = re.search(
        r"(?i)verdict\s*[:\-]\s*\**\s*(conditional go|no-go|no go|pivot|go)\b",
        report_text or "",
    )
    return match.group(1).title() if match else None


def _normalise_report(report_text, startup_name):
    """Add a title if the model omitted one and remove surrounding code fences."""
    text = str(report_text or "").strip()
    text = re.sub(r"^\s*```(?:markdown|md|text)?\s*", "", text, flags=re.I)
    text = re.sub(r"\s*```\s*$", "", text)
    text = _clean_latex(text)
    if not re.search(r"(?im)^#\s+", text):
        text = f"# Executive Assessment Report: {startup_name}\n\n{text}"
    return text.strip()


def _markdown_inline(paragraph, text):
    """Render a small subset of Markdown inline formatting into a DOCX paragraph."""
    parts = re.split(r"(\*\*.*?\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = DocxPt(9)
        else:
            paragraph.add_run(part)


def _add_md_table(doc, table_lines):
    """Turn Markdown table lines into a real Word table."""
    rows = []
    for ln in table_lines:
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
            continue  # separator row
        rows.append(cells)
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            _markdown_inline(p, row[j] if j < len(row) else "")
            for run in p.runs:
                run.font.size = DocxPt(9)
                if i == 0:
                    run.bold = True
    doc.add_paragraph()


def generate_docx(report_text, filename="PitchWise_Executive_Report.docx",
                  startup_name="Startup Assessment"):
    """Create a polished, editable Word executive assessment report."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUTPUT_DIR / filename
    report_text = _normalise_report(report_text, startup_name)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = DocxInches(0.7)
    section.bottom_margin = DocxInches(0.65)
    section.left_margin = DocxInches(0.8)
    section.right_margin = DocxInches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = DocxPt(10)
    styles["Normal"].font.color.rgb = DocxRGBColor(45, 55, 72)
    styles["Normal"].paragraph_format.space_after = DocxPt(5)

    for style_name, size, color in [
        ("Title", 26, DocxRGBColor(22, 37, 66)),
        ("Heading 1", 17, DocxRGBColor(22, 37, 66)),
        ("Heading 2", 13, DocxRGBColor(48, 95, 170)),
        ("Heading 3", 11, DocxRGBColor(25, 139, 145)),
    ]:
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = DocxPt(size)
        style.font.bold = True
        style.font.color.rgb = color

    # Header and footer
    header = section.header.paragraphs[0]
    header.text = f"PITCHWISE  |  EXECUTIVE ASSESSMENT  |  {startup_name.upper()}"
    header.style = styles["Caption"]
    for run in header.runs:
        run.font.color.rgb = DocxRGBColor(103, 114, 130)
        run.font.size = DocxPt(8)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Prepared by PitchWise • AI-assisted assessment • Validate material claims before acting").font.size = DocxPt(8)
    for run in footer.runs:
        run.font.color.rgb = DocxRGBColor(103, 114, 130)

    lines = report_text.splitlines()
    first_heading_used = False
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue

        # Markdown table: collect consecutive lines starting with "|"
        if line.startswith("|"):
            table_lines = [line]
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            _add_md_table(doc, table_lines)
            continue

        if line.startswith("# "):
            title = line[2:].strip()
            if not first_heading_used:
                p = doc.add_paragraph(style="Title")
                p.paragraph_format.space_after = DocxPt(8)
                _markdown_inline(p, title)
                first_heading_used = True
            else:
                p = doc.add_paragraph(style="Heading 1")
                _markdown_inline(p, title)
        elif line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            _markdown_inline(p, line[3:].strip())
        elif line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            _markdown_inline(p, line[4:].strip())
        elif re.match(r"^\s*[-*]\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            _markdown_inline(p, re.sub(r"^\s*[-*]\s+", "", line))
        elif re.match(r"^\s*\d+[.)]\s+", line):
            p = doc.add_paragraph(style="List Number")
            _markdown_inline(p, re.sub(r"^\s*\d+[.)]\s+", "", line))
        elif line.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = DocxInches(0.25)
            p.paragraph_format.space_before = DocxPt(3)
            p.paragraph_format.space_after = DocxPt(7)
            _markdown_inline(p, line.lstrip("> "))
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = DocxRGBColor(80, 93, 112)
        else:
            p = doc.add_paragraph()
            _markdown_inline(p, line)

    doc.save(path)
    return str(path)


def _parse_slide_outline(slide_outline):
    """Parse the designer agent's 'Slide N: Title' and '- bullet' format."""
    slides = []
    current = None

    for raw in _clean_latex(slide_outline).splitlines():
        line = raw.strip()
        match = re.match(r"(?i)^slide\s*(\d+)\s*:\s*(.+)$", line)
        if match:
            if current:
                slides.append(current)
            current = {"number": int(match.group(1)), "title": match.group(2).strip(), "bullets": []}
            continue

        if current and re.match(r"^[-*•]\s+", line):
            bullet = re.sub(r"^[-*•]\s+", "", line).strip()
            if bullet:
                current["bullets"].append(_strip_markdown(bullet))

    if current:
        slides.append(current)

    slides = [s for s in slides if s["title"]]
    for slide in slides:
        slide["bullets"] = slide["bullets"][:5]
    return slides


def _add_text(slide, text, left, top, width, height, font_size=18,
              color=INK, bold=False, font_name="Aptos",
              align=PP_ALIGN.LEFT, margin=0.06):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = str(text)
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def _add_footer(slide, startup_name, slide_num, total):
    _add_text(slide, startup_name, 0.55, 7.12, 5.2, 0.18, 8, MUTED)
    _add_text(slide, f"PITCHWISE  •  {slide_num}/{total}", 9.0, 7.12, 3.75, 0.18,
              8, MUTED, align=PP_ALIGN.RIGHT)


def _add_cover(prs, startup_name, verdict_hint="Executive Assessment"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY

    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.22), Inches(7.5)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL
    accent.line.fill.background()

    _add_text(slide, "PITCHWISE  /  STRATEGIC REVIEW", 0.85, 1.25, 7.0, 0.35,
              13, RGBColor(135, 215, 211), True)
    _add_text(slide, startup_name, 0.85, 2.0, 11.0, 1.3, 34, WHITE, True,
              font_name="Aptos Display")
    _add_text(slide, verdict_hint, 0.9, 3.5, 9.8, 0.55, 21,
              RGBColor(218, 226, 239), False)
    _add_text(slide, "Executive assessment and recommended next steps",
              0.9, 5.7, 10.5, 0.4, 14, RGBColor(218, 226, 239))
    _add_text(slide, datetime.now().strftime("%d %B %Y"), 0.9, 6.35, 4.0, 0.3,
              10, RGBColor(174, 188, 208))
    return slide


def generate_pptx(slide_outline, filename="PitchWise_Client_Presentation.pptx",
                  startup_name="Startup Assessment", report_text=""):
    """Create an editable widescreen PowerPoint from the slide outline."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUTPUT_DIR / filename
    content_slides = _parse_slide_outline(slide_outline)

    if not content_slides:
        raise ValueError(
            "No slides were parsed from the Presentation Designer output. "
            "Expected headings such as 'Slide 1: Startup Overview' followed by '- bullets'."
        )

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    verdict = _extract_verdict(report_text)
    verdict_hint = (f"Executive Assessment  |  Preliminary verdict: {verdict}"
                    if verdict else "Executive Assessment")
    _add_cover(prs, startup_name, verdict_hint)

    total = len(content_slides)
    for idx, item in enumerate(content_slides, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = WHITE

        band = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.16)
        )
        band.fill.solid()
        band.fill.fore_color.rgb = TEAL if idx in (1, total) else BLUE
        band.line.fill.background()

        _add_text(slide, f"{idx:02d}", 0.55, 0.48, 0.55, 0.4, 14, TEAL, True)
        _add_text(slide, item["title"], 1.15, 0.42, 11.45, 0.62,
                  27, NAVY, True, font_name="Aptos Display")

        divider = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.18), Inches(12.1), Inches(0.025)
        )
        divider.fill.solid()
        divider.fill.fore_color.rgb = RGBColor(221, 228, 237)
        divider.line.fill.background()

        bullets = item["bullets"] or ["No slide points were returned for this section."]
        available_height = 5.45
        row_height = min(1.18, available_height / max(len(bullets), 1))
        start_y = 1.48

        for bullet_index, bullet in enumerate(bullets):
            y = start_y + bullet_index * row_height

            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(0.68), Inches(y), Inches(11.98), Inches(row_height - 0.12)
            )
            card.fill.solid()
            card.fill.fore_color.rgb = PALE
            card.line.color.rgb = RGBColor(231, 236, 243)

            dot = slide.shapes.add_shape(
                MSO_SHAPE.OVAL, Inches(0.92), Inches(y + 0.22), Inches(0.12), Inches(0.12)
            )
            dot.fill.solid()
            dot.fill.fore_color.rgb = TEAL if bullet_index == 0 else BLUE
            dot.line.fill.background()

            clean_bullet = _strip_markdown(bullet)
            font_size = 16 if len(clean_bullet) < 145 else 14
            _add_text(
                slide, clean_bullet, 1.22, y + 0.09, 11.0, row_height - 0.23,
                font_size, INK, False, margin=0.03
            )

        _add_footer(slide, startup_name, idx, total)

    prs.save(path)
    return str(path)


def _verbose():
    """Set PITCHWISE_VERBOSE=1 to see CrewAI's live agent logs in the console."""
    return os.environ.get("PITCHWISE_VERBOSE") == "1"


def _make_agent(role, goal, backstory, llm):
    return Agent(
        role=role,
        goal=goal,
        backstory=backstory,
        llm=llm,
        verbose=_verbose(),
        allow_delegation=False,
    )


async def run_pitchwise(
    pitch,
    api_key=None,
    model=None,
    report_filename=None,
    pptx_filename=None,
    output_dir=None,
    verbose_print=True,
    on_progress=None,
    generate_files=True,
    skip_critic=False,
):
    """
    Run all six agents and generate a DOCX report plus PPTX presentation.

    In a Colab/Jupyter cell use:   result = await run_pitchwise(pitch)
    In a script or Streamlit use:  result = run_pitchwise_sync(pitch)

    Returns a dictionary:
      {
        "startup_name": ..., "report": ..., "slide_outline": ...,
        "steps": [(agent_role, output_text), ...],
        "docx_path": "...", "pptx_path": "..."
      }
    """
    global OUTPUT_DIR
    if output_dir:
        OUTPUT_DIR = Path(output_dir).resolve()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    startup_name = _extract_startup_name(pitch)
    safe_name = _safe_filename(startup_name)
    report_filename = report_filename or f"{safe_name}_Executive_Report.docx"
    pptx_filename = pptx_filename or f"{safe_name}_Client_Presentation.pptx"

    llm = get_llm(api_key=api_key, model=model)

    market_analyst = _make_agent(
        role="Market Demand Analyst",
        goal="Assess customer demand, target segments, willingness to pay, and what evidence is missing.",
        backstory=(
            "You are a skeptical market strategy analyst. Distinguish pitch-provided facts from "
            "assumptions. Do not invent customer interviews, market sizes, LOIs, pilot results, or sources."
        ),
        llm=llm,
    )
    competitive_analyst = _make_agent(
        role="Competitive Analyst",
        goal="Identify direct and indirect competitors, substitutes, positioning gaps, and differentiation risks.",
        backstory=(
            "You are a competitive strategy analyst. Separate known competitor categories from claims "
            "that require current external research. Do not claim that a tool is free or has a specific capability "
            "unless the supplied information establishes it."
        ),
        llm=llm,
    )
    financial_analyst = _make_agent(
        role="Financial Analyst",
        goal="Evaluate pricing, revenue assumptions, delivery costs, acquisition targets, and missing financial inputs.",
        backstory=(
            "You are a commercially rigorous financial analyst. Do not fabricate financial figures. "
            "You show the arithmetic behind every number you compute and double-check that totals equal "
            "the sum of their parts. Explain what can and cannot be concluded from the pitch."
        ),
        llm=llm,
    )
    risk_critic = _make_agent(
        role="Risk and Assumption Critic",
        goal="Rank the most important risks and propose concrete validation or mitigation steps.",
        backstory=(
            "You are an independent risk reviewer. Avoid alarmist language and unsupported absolutes. "
            "Rank risks by impact and uncertainty, and propose practical tests."
        ),
        llm=llm,
    )
    strategy_writer = _make_agent(
        role="Strategy Lead and Report Writer",
        goal="Synthesize the analysis into a clear executive assessment and actionable recommendations.",
        backstory=(
            "You are a management consultant writing for a founder, executive, or client. Be concise, "
            "professional, specific, and candid. Distinguish evidence, assumptions, and recommendations. "
            "Your verdict must be consistent with the risks and recommendations you write."
        ),
        llm=llm,
    )
    presentation_designer = _make_agent(
        role="Presentation Designer",
        goal="Convert the final assessment into a concise, client-facing slide outline.",
        backstory=(
            "You are an executive presentation designer. Use short, decision-oriented bullets, not paragraphs. "
            "Do not introduce claims that conflict with the supporting analysis."
        ),
        llm=llm,
    )

    t1 = Task(
        description=f"""
Analyze the following startup pitch.

PITCH:
{pitch}

Assess:
1. Customer problem and likely demand drivers.
2. Most promising initial customer segment and why.
3. Willingness-to-pay assumptions and evidence gaps.
4. Claims that are not yet validated.
5. Three practical customer or pilot validations.

Clearly label what is explicitly stated in the pitch versus inferred. Do not invent research, customer evidence, market-size figures, or citations.
Length limit: 250 words.
{COMMON_RULES}
""",
        expected_output="A structured market demand assessment (max 250 words) with evidence gaps and validation steps.",
        agent=market_analyst,
    )

    t2 = Task(
        description=f"""
Analyze the competitive position for this startup.

PITCH:
{pitch}

Identify:
1. Direct competitor categories and examples, where known.
2. Indirect competitors and substitutes.
3. Likely differentiation and defensibility gaps.
4. A realistic positioning opportunity.
5. What must be researched or verified before presenting the analysis as fact.

Do not imply that you performed live research. Mark potentially time-sensitive competitor details as requiring verification.
Length limit: 250 words.
{COMMON_RULES}
""",
        expected_output="A structured competitor assessment (max 250 words) that separates plausible analysis from verified evidence.",
        agent=competitive_analyst,
    )

    t3 = Task(
        description=f"""
Evaluate the business model and financial assumptions in this pitch.

PITCH:
{pitch}

Assess pricing, recurring revenue, sales assumptions, acquisition costs, delivery costs, gross margin drivers, and the profitability claim. List the specific missing inputs needed for a defensible financial model. Do not fabricate numbers or call the business financially viable/non-viable without explaining the evidence.

Arithmetic rules: for every figure you compute, write the formula (for example "20 x Rs 25 lakh = Rs 500 lakh = Rs 5 crore"). Before finishing, re-add any totals and confirm they equal the sum of their components. State any timing assumption (such as clients arriving gradually) instead of assuming everything happens on day one.
Length limit: 300 words.
{COMMON_RULES}
""",
        expected_output="A financial assessment (max 300 words) with shown arithmetic, explicit assumptions, risks, and missing inputs.",
        agent=financial_analyst,
    )

    t4 = Task(
        description=f"""
Use the market, competitive, and financial analyses to rank the five most important risks or assumptions.
For each risk, provide:
- Severity (Critical / High / Medium / Low)
- Why it matters (1-2 sentences)
- What evidence is currently missing
- A concrete test or mitigation
Avoid repeating the same risk in different words. Avoid sensational or absolute claims unless directly justified.
Length limit: 300 words.
{COMMON_RULES}
""",
        expected_output="Five prioritized risks (max 300 words), each with severity, rationale, evidence gap, and action.",
        agent=risk_critic,
        context=[t1, t2, t3],
    )

    t5 = Task(
        description=f"""
Write the authoritative executive assessment report for this pitch.

PITCH:
{pitch}

Synthesize the preceding analyses into a professional report of 700-1,000 words. Use Markdown headings in this order:
# Executive Assessment: [Startup Name]
## 1. Executive Summary
## 2. Business and Customer Opportunity
## 3. Market Demand and Target Segment
## 4. Competitive Landscape and Differentiation
## 5. Business Model and Financial Assessment
## 6. Key Risks and Assumptions
## 7. Prioritized Recommendations
## 8. Verdict and Immediate Next Steps
## Evidence Limitations

Verdict rules:
- Put one line "Verdict: <choice>" near the beginning and explain it in Section 8.
- Choose exactly one: Go, Conditional Go, Pivot, or No-Go.
- Go = no Critical risks and key claims are supported.
- Conditional Go = the core model is sound and the open risks can be settled by cheap validation tests.
- Pivot = the pricing model, target segment, or product scope must change before proceeding.
- No-Go = a fundamental flaw that validation is unlikely to fix.
- Your verdict must match your own recommendations: if you recommend changing the pricing model or target segment, or if two or more risks are Critical, the verdict must be Pivot or No-Go, not Conditional Go.

Content rules:
- Separate pitch-provided facts from unverified founder claims, analysis, and recommendations.
- Do not invent statistics, interviews, LOIs, pilot outcomes, financial inputs, or citations.
- Avoid presenting market or procurement-cycle statements as verified facts unless supported by supplied evidence.
- Include 3-5 prioritized recommendations with a practical validation measure or deliverable for each. Describe each validation test once, in Section 7, and refer to it elsewhere instead of repeating it.
- Reuse the financial arithmetic from the Financial Analyst exactly; do not recompute with different numbers.
- Use professional, client-appropriate language. Avoid inflammatory phrases such as "delusion" or "baseless".
- State clearly that the assessment is based on the pitch and agent analysis, not independent market research, unless sources were supplied.
{COMMON_RULES}
""",
        expected_output="A client-ready Markdown executive assessment report (700-1,000 words) with a consistent verdict and explicit evidence limitations.",
        agent=strategy_writer,
        context=[t1, t2, t3] + ([] if skip_critic else [t4]),
    )

    t6 = Task(
        description=f"""
Create exactly seven content slides for a client-facing presentation, based on the pitch and the executive assessment report.

Use this exact format:
Slide 1: Startup Overview
- concise bullet
- concise bullet

Slide 2: Problem and Proposed Solution
- concise bullet

Slide 3: Market Demand and Target Customers
- concise bullet

Slide 4: Competitive Landscape
- concise bullet

Slide 5: Business Model and Financial Assessment
- concise bullet

Slide 6: Key Risks and Validation Priorities
- concise bullet

Slide 7: Recommendations and Verdict
- concise bullet

Requirements:
- Provide 3-4 concise bullets per slide.
- Keep each bullet ideally under 20 words; avoid paragraph-length bullets.
- Use only figures present in the pitch or analysis, and label founder claims as unverified where appropriate.
- Ensure the verdict exactly matches the executive assessment report.
- No opening commentary, Markdown tables, invented citations, or text outside the slide outline.
{COMMON_RULES}
""",
        expected_output="Exactly seven slide headings, each followed by 3-4 concise, decision-oriented bullets.",
        agent=presentation_designer,
        context=[t1, t2, t3] + ([] if skip_critic else [t4]) + [t5],
    )

    pipeline = [
        ("Market Demand Analyst", market_analyst, t1),
        ("Competitive Analyst", competitive_analyst, t2),
        ("Financial Analyst", financial_analyst, t3),
    ]
    if not skip_critic:
        pipeline.append(("Risk and Assumption Critic", risk_critic, t4))
    pipeline += [
        ("Strategy Lead and Report Writer", strategy_writer, t5),
        ("Presentation Designer", presentation_designer, t6),
    ]
    agent_names = [p[0] for p in pipeline]
    tasks = [p[2] for p in pipeline]
    total_steps = len(tasks)
    counter = {"n": 0}

    def _task_cb(_output):
        counter["n"] += 1
        idx = min(counter["n"], total_steps) - 1
        if on_progress:
            try:
                on_progress(counter["n"], total_steps, agent_names[idx])
            except Exception:
                pass
        elif verbose_print:
            print(f"[{counter['n']}/{total_steps}] {agent_names[idx]} finished")

    max_rpm = int(os.environ.get("PITCHWISE_MAX_RPM", "0")) or None  # set e.g. 5 on free tiers

    crew = Crew(
        agents=[p[1] for p in pipeline],
        tasks=tasks,
        process=Process.sequential,
        max_rpm=max_rpm,
        task_callback=_task_cb,
        verbose=_verbose(),
    )

    # Run the agents ONCE
    crew_output = await crew.kickoff_async()

    outputs = getattr(crew_output, "tasks_output", None) or [t.output for t in tasks]

    formatted_steps = []
    for name, out in zip(agent_names, outputs):
        text = _clean_latex(out.raw if hasattr(out, "raw") else str(out))
        formatted_steps.append((name, text))

    steps_by_name = dict(formatted_steps)
    report_text = steps_by_name.get("Strategy Lead and Report Writer", "")
    slide_outline = steps_by_name.get("Presentation Designer", "")

    if verbose_print:
        print("\n" + "=" * 90)
        print(f"PITCHWISE MULTI-AGENT ANALYSIS: {startup_name}")
        print("=" * 90)
        for index, (name, text) in enumerate(formatted_steps, start=1):
            print(f"\n{'=' * 90}")
            print(f"AGENT {index}: {name.upper()}")
            print("=" * 90)
            print(text)

    docx_path = pptx_path = None
    if generate_files:
        docx_path = generate_docx(
            report_text,
            filename=report_filename,
            startup_name=startup_name,
        )
        pptx_path = generate_pptx(
            slide_outline,
            filename=pptx_filename,
            startup_name=startup_name,
            report_text=report_text,
        )

    if verbose_print and generate_files:
        print("\n" + "=" * 90)
        print("DELIVERABLES GENERATED")
        print("=" * 90)
        print("Executive report:", docx_path)
        print("Client presentation:", pptx_path)

    return {
        "startup_name": startup_name,
        "verdict": _extract_verdict(report_text),
        "report": report_text,
        "slide_outline": slide_outline,
        "steps": formatted_steps,
        "docx_path": docx_path,
        "pptx_path": pptx_path,
    }


def run_pitchwise_sync(pitch, **kwargs):
    """Synchronous wrapper for scripts, Streamlit and evaluate.py.
    Also safe inside Colab, where an event loop is already running."""
    try:
        asyncio.get_running_loop()
    except RuntimeError:
        return asyncio.run(run_pitchwise(pitch, **kwargs))

    box = {}

    def _target():
        try:
            box["value"] = asyncio.run(run_pitchwise(pitch, **kwargs))
        except Exception as e:  # re-raised in the caller's thread
            box["error"] = e

    thread = threading.Thread(target=_target)
    thread.start()
    thread.join()
    if "error" in box:
        raise box["error"]
    return box["value"]
